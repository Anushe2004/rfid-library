"""Read .xlsx / .csv sources and map headers to canonical fields."""

import csv
import json
from pathlib import Path

FIELDS = ["accession_no", "barcode", "isbn", "title", "author", "publisher", "pub_year"]


class SchemaError(Exception):
    pass


def load_mapping(path) -> dict:
    return json.loads(Path(path).read_text())


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _raw_rows(source: Path):
    if source.suffix.lower() == ".csv":
        with source.open(newline="", encoding="utf-8-sig") as fh:
            yield from csv.reader(fh)
    elif source.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(source, read_only=True, data_only=True)
        yield from wb.worksheets[0].iter_rows(values_only=True)
        wb.close()
    else:
        raise SchemaError(f"Unsupported file type: {source.suffix}")


def read_rows(source, mapping: dict):
    """Yield (row_no, canonical_dict). row_no is the 1-based spreadsheet row (header = 1)."""
    it = _raw_rows(Path(source))
    header = next(it, None)
    if header is None:
        raise SchemaError("Source file is empty")
    norm = {_cell(h).lower(): i for i, h in enumerate(header)}
    index = {}
    for field, aliases in mapping["columns"].items():
        for a in aliases:
            if a.lower() in norm:
                index[field] = norm[a.lower()]
                break
    missing = [f for f in mapping["required"] if f not in index]
    if missing:
        raise SchemaError(f"Required column(s) not found in source: {', '.join(missing)}")
    for n, row in enumerate(it, start=2):
        if row is None or all(_cell(c) == "" for c in row):
            continue
        yield n, {f: (_cell(row[i]) if i < len(row) else "") for f, i in index.items()}
